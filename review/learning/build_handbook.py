"""Build a read-only, evidence-linked Chinese handbook for the accepted model."""
from pathlib import Path
import json, re, hashlib, collections, shutil, unicodedata
import numpy as np
from PIL import Image, ImageChops
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
MODEL=Path('D:/simulink/modern/PMSM_FOC_R2026a.slx')
EXPECTED='a745187d8a8b40986515717c7ebc6594f488e745b7f73bd922bf1bdc39b46994'
assert hashlib.sha256(MODEL.read_bytes()).hexdigest()==EXPECTED
E=json.loads((ROOT/'current_model_evidence.json').read_text(encoding='utf-8'))
B=E['blocks']; M='PMSM_FOC_R2026a'; ASSET=ROOT/'handbook_assets'; ASSET.mkdir(exist_ok=True)
def arr(x): return [x] if isinstance(x,dict) else x
def clean(x): return str(x).replace('\n',' ').replace('\r',' ').strip()
def sid(b): return b['sid'].removeprefix(M+':')
BY={b['path']:b for b in B}; BS={sid(b):b for b in B}
G={b['params']['GotoTag']:b for b in B if b['type']=='Goto'}
F=collections.defaultdict(list)
for b in B:
 if b['type']=='From': F[b['params']['GotoTag']].append(b)
groups=[
[592,169,155,193,222,509,201,190,196,216,578,217,215,195,194,197],
[590,40,439,144,220,202,203,206,210,139,141,212,213,579,133,453,152,221,209,214,468],
[200,183,593,597,62,65,66,265,199,191,228,229,232,231,606,236,237,238,266,233,235],
[242,243,246,247,240,295,346,347,348,291,292,293,299,300,301,298,244,324,426,302,307,372,718],
[681,245,615,671,626,523,682,496,506,281,693],
[649,32,28,34,171,172,186,188,177,180,189,660,358,359],
[259,271,255,310,312,314,273,274,282,334,288,361,360,365,362,363,366],
[269,275,285,286,287,350,352,353,331,332,330,308,309,497,498,499],
[501,589,534,504,546,548,551,550,549,586,576,563,564,566,558,559,560,565,683,553,588,688,573,575,695]]
assert len(set(sum(groups,[])))==164
assert set(map(str,sum(groups,[])))==set(sid(b) for b in B if b['parent']==M)
tags={'speed':'机械转速，rpm','nref':'主转速给定，rpm','nref2':'辅助速度 Chirp，rpm',
'id':'d 轴电流，A','iq':'q 轴电流，A','iq_ref':'主 q 轴给定，A','id_ref':'辅助 d 轴 Chirp，A',
'ws':'电角速度，rad/s','theta_e':'电角度，rad','vd':'最终 d 轴电压要求，V','vq':'最终 q 轴电压要求，V',
'v_alpha':'α 轴电压要求，V','v_beta':'β 轴电压要求，V','Te':'电磁转矩，N·m','TL':'负载转矩，N·m',
'UDC':'理想母线测量，V','SW':'六路门极 [AH AL BH BL CH CL]',
'Iabc':'三相电流观测向量，A','Currentabc':'三相电流控制反馈，A',
'Uabc':'线电压 [Vab Vbc Vca]，V','Vq_re':'电机端实际 q 轴电压，V','Vd_re':'电机端实际 d 轴电压，V',
'Umod':'三相比较时间 Tcm，s','Tcm':'三相比较时间 Tcm，s','sector':'扇区编码，取 0 至 6'}
roles={
592:'主转速给定 n=3000 rpm；常量直接进入速度误差。',169:'目标转速减去保持后的实际转速，得到 rpm 误差。',
155:'速度 PI；输出按 Te_ref 使用，增益单位及面板采样参数待统一。',193:'转矩要求除以当前换算分母，输出 iq_ref，A。',
196:'永磁磁链 phis=0.0054 Wb。',216:'d 轴电感 Ld=120 μH。',217:'计算 Ld·id，当前转矩换算的一部分。',
215:'计算 phis+Ld·id；它与完整凸极转矩分母有差别。',195:'极对数 p=4。',194:'幅值不变变换的转矩系数 1.5。',197:'计算 1.5·p·(phis+Ld·id)，作为电流换算分母。',
590:'主 d 轴给定为常量 0 A；与辅助标签 id_ref 独立。',40:'id_ref−id，得到 d 轴电流误差。',
439:'d 轴 PI，但当前遮罩引用 q 轴增益和限幅；内部周期为 dt。',144:'d 轴 PI 减去 ωe·Lq·iq，输出最终 vd。',
139:'q 轴电感 Lq=120 μH。',141:'计算 ωe·Lq·iq，d 轴交叉耦合前馈。',212:'永磁磁链 phis。',
213:'计算 ωe·phis，当前 q 轴反电动势前馈；不含 Ld·id。',133:'端口 1 减 iq、端口 2 加 iq_ref，结果为 iq_ref−iq。',
453:'q 轴离散 PI，含积分状态限幅和饱和误差反馈。',152:'q 轴 PI 与 ωe·phis 相加，输出最终 vq。',
209:'d 轴电感 Ld。',214:'计算 Ld·id，但输出未接入 q 轴前馈。',468:'终止未使用的 Ld·id 支路，不参与主控制计算。',
593:'Clarke 变换，幅值不变；输出 [α β 0]。',597:'Park 变换，使用采样电角度得到 [id iq i0]。',
62:'拆分 Park 输出；依次为 id、iq、i0。',265:'终止当前不使用的电流零序分量。',
232:'逆变换零序电压给定为 0 V。',231:'按 [vd vq 0] 顺序组合输入。',606:'逆 Park 变换；输出 [vα vβ v0]。',
236:'拆分逆 Park 输出的三个坐标分量。',266:'终止不使用的电压零序分量。',
246:'调制使用的母线 Vdc=24 V，当前为常量。',247:'调制周期 dt=50 μs。',
240:'根据 αβ 电压计算 Tcm 与 sector；完整源码和现有问题见正文及附录。',295:'将 Tcm 向量拆成 A、B、C 三相比较阈值。',
324:'周期三角载波，时间点 [0 Ts/2 Ts]，幅值 [0 Ts/2 0]，单位 s。',
291:'载波≥A 相阈值时上管为 1。',292:'载波≥B 相阈值时上管为 1。',293:'载波≥C 相阈值时上管为 1。',
299:'A 相上管逻辑取反生成下管；无死区。',300:'B 相上管逻辑取反生成下管；无死区。',301:'C 相上管逻辑取反生成下管；无死区。',
298:'按 [AH AL BH BL CH CL] 组合六路门极。',302:'四端口依次观察载波与 A、B、C 三相阈值。',
681:'理想恒定直流电源 Vdc=24 V；没有源阻抗或电容状态。',615:'理想六开关桥计算浮动星点三相电压，未建模器件和续流。',
671:'输入相电压和相电流，输出线电压、电流反馈和传给电机的相电压。',
626:'连续 dq 电机方程；四个状态为 id、iq、ωm、θm，另计算测量向量和 Iabc。',
523:'负载 Step：Time=0，前后均为 0.3·Tn=0.089424 N·m，SampleTime=dt。',
682:'Gain=1，原样发布电源数值作为母线测量。',
649:'拆分八元素电机测量向量，输出前六项；不是电气总线块。',32:'将 ia、ib、ic 重新组合成电流观测向量。',
28:'三个输入端口分别观察 ia、ib、ic。',171:'30/π，将机械 rad/s 转为 rpm。',186:'乘极对数 p，将机械角速度转为电角速度。',
177:'θm 乘 p 后 mod 2π，生成理想电角度。',660:'从同一测量向量只取第 7 项 vq_actual 与第 8 项 vd_actual。',
255:'端口顺序 speed、Te；分别读 rpm 和 N·m。',314:'端口顺序 id、iq，不包含 iq_ref。',274:'观察 Iabc 三元素向量。',
334:'拆开 Uabc 的 [Vab Vbc Vca]，不是三相相电压。',288:'依次观察三相线电压 Vab、Vbc、Vca。',
365:'端口顺序 Vd_re、Vq_re，即电机实际 dq 电压。',366:'端口顺序 vd、vq，即控制器电压要求。',
275:'观察六路 SW 门极向量。',287:'端口顺序 v_alpha、v_beta。',352:'拆开 Umod，也就是三相 Tcm。',
353:'依次观察三相 Tcm，单位 s。',330:'横轴 v_alpha、纵轴 v_beta，观察电压矢量轨迹。',
309:'观察扇区编码；物理顺序与数字排序不同。',499:'端口顺序 UDC、speed，单位分别为 V 和 rpm。',
534:'组合实际 speed 与辅助 nref2。',504:'端口 1 是双信号向量，端口 2 是归一化绝对差值。',
546:'辅助 nref2 减去实际 speed；不是主速度误差。',548:'速度辅助差值取绝对值。',
551:'max(abs(0.1*n),1)=300 rpm，速度辅助误差分母。',550:'速度绝对差值除以 300 rpm，得到无量纲量。',549:'端口顺序实际 speed、辅助 nref2。',
563:'组合实际 id 与辅助 Chirp 标签 id_ref。',564:'端口 1 是双信号向量，端口 2 是归一化绝对差值。',
566:'辅助 id_ref 减去实际 id；不是主 d 轴给定误差。',558:'辅助 d 轴差值取绝对值。',559:'0.1*In=0.92 A，辅助误差分母。',
560:'辅助电流绝对差值除以 0.92 A。',565:'端口顺序实际 id、辅助 id_ref。',
683:'10 s 内由 0 扫至 500 Hz 的余弦 Chirp，以 dt 保持；不驱动主给定。',
553:'0.1*n=300，将单位 Chirp 缩放为辅助 rpm。',
688:'10 s 内由 0 扫至 5000 Hz 的余弦 Chirp，以 dt 保持；不驱动主给定。',573:'0.1*In=0.92，将 Chirp 缩放为辅助 A。',
695:'只读 12 个标签并按 dt 存 timeseries；不参与控制反馈。'}
def role(b):
 s=sid(b); t=b['type']; p=b['params']; n=clean(b['name']); parent=clean(b['parent'].removeprefix(M+'/'))
 if b['parent']==M and s.isdigit() and int(s) in roles:return roles[int(s)]
 if t in ('From','Goto'):
  tag=p['GotoTag']; assert tag in tags,tag
  return ('读取' if t=='From' else '发布')+'标签 '+tag+'：'+tags[tag]+'；不改变数值。'
 if t=='ZeroOrderHold':return '按 '+p['SampleTime']+'=50 μs 采样保持；两次更新之间输出保持不变。'
 if t=='RateTransition':return '采样率过渡，OutPortSampleTime='+p.get('OutPortSampleTime','继承')+'；数据完整性开启、确定性关闭；不能等同完整 PWM 延迟。'
 if parent.startswith('PI Controller'):
  d={'error':'子系统误差入口；速度为 rpm，电流环为 A。','Kp':'误差乘遮罩 Kp，生成比例项。',
  'Ki':'抗饱和修正后的误差乘遮罩 Ki，驱动积分器。','Gain':'限幅前后差乘 Ks；后面还会乘 Ki。',
  'Gain1':'误差乘 Bm；当前为 0，并非机械摩擦补偿。','Sum':'计算限幅前输出减去限幅后输出。',
  'Sum1':'从原误差减去 Ks 倍饱和差。','Add':'比例项加积分项'+('再减 Bm·error。' if parent=='PI Controller_w' else '。'),
  'Saturation':'总输出限制在 downlimit 与 uplimit。','Scope2':'观察比例支路保持后的值。','Scope3':'观察积分器输出。','Out':'向外输出限幅后的 PI 结果。'}
  if n in d:return d[n]
  if t=='DiscreteIntegrator':return '后向欧拉离散积分，SampleTime=dt、初值 0，积分状态限制在 downlimit 至 uplimit。'
 if parent=='Permanent Magnet Synchronous Machine':
  d={'Vabc':'端口 1 输入相电压 [va vb vc]，V。','load_torque':'端口 2 输入 TL，N·m。',
  'id_iq_wm_theta':'连续积分四个状态，初值 [0;0;0;0]；导数顺序必须一致。',
  'v_state_load':'组合 [va vb vc id iq wm θm TL] 八元素计算输入。',
  'derivatives':'组合 [did/dt diq/dt dωm/dt dθm/dt]。',
  'did_dt':'按 d 轴电压方程计算 (vd−Rs·id+ωe·Lq·iq)/Ld，A/s。',
  'diq_dt':'按 q 轴电压方程计算 (vq−Rs·iq−ωe·(Ld·id+phis))/Lq，A/s。',
  'dwm_dt':'由电磁转矩减负载和黏性摩擦再除惯量，得到机械角加速度。',
  'dtheta_dt':'取 ωm 作为机械角度的导数。',
  'measurement':'组合 [ia ib ic ωm θm Te vq_actual vd_actual]。',
  'ia':'用 θe=p·θm 将 dq 电流变回 a 相。','ib':'以 θe−2π/3 将 dq 电流变回 b 相。','ic':'以 θe+2π/3 将 dq 电流变回 c 相。',
  'wm':'从计算向量第 6 项取机械角速度，rad/s。','theta_m':'取第 7 项机械角，rad。',
  'Te':'计算 1.5·p·[phis·iq+(Ld−Lq)·id·iq]，N·m。',
  'vq_actual':'按电机角度将实际相电压变换为 q 轴电压，V。','vd_actual':'按电机角度将实际相电压变换为 d 轴电压，V。',
  'measurement_vector':'端口 1 输出八元素测量向量。','current_vector':'组合三相电流 [ia ib ic]。','Iabc':'端口 2 输出 Iabc，A。'}
  assert n in d,(parent,n);return d[n]
 if parent in ('Alpha-Beta-Zero to dq0','dq0 to Alpha-Beta-Zero'):
  inv=parent.startswith('dq0');d={
  'alpha_beta_zero':('输出 [α β 0]。' if inv else '端口 1 输入 [α β 0]。'),
  'dq0':('端口 1 输入 [d q 0]。' if inv else '输出 [d q 0]。'),
  'theta_e':'端口 2 输入电角度，rad。','input_vector':'组合坐标三分量与 θe；第 4 项为角度。',
  'axis_1':('计算 d·cosθe−q·sinθe，得到 α。' if inv else '计算 α·cosθe+β·sinθe，得到 d。'),
  'axis_2':('计算 d·sinθe+q·cosθe，得到 β。' if inv else '计算 −α·sinθe+β·cosθe，得到 q。'),
  'axis_0':'原样取第三项零序分量。','result':'组合三个坐标分量，顺序由子系统输出约定。'}
  assert n in d,(parent,n);return d[n]
 if parent=='abc to Alpha-Beta-Zero':
  d={'abc':'输入 [ia ib ic]，A。','Clarke matrix':'乘幅值不变 Clarke 矩阵，含零序；矩阵与公式见第 03 区。','alpha_beta_zero':'输出 [iα iβ i0]，A。'}
  assert n in d,(parent,n);return d[n]
 if parent=='Universal Bridge':
  d={'six_gates':'输入 [AH AL BH BL CH CL] 六路逻辑。','Vdc':'输入母线电压，V。','double gates':'门极转换为 double，供算术表达式使用。',
  'gates_and_bus':'把六路门极和 Vdc 组合为七元素向量。','phase_A':'按浮动中性点公式计算 a 相电压，见第 05 区。',
  'phase_B':'按相同公式计算 b 相电压。','phase_C':'按相同公式计算 c 相电压。','phase_voltage':'组合 [va vb vc]。','Vabc':'输出电机相电压 [va vb vc]，V。'}
  assert n in d,(parent,n);return d[n]
 if parent=='Three-Phase V-I Measurement':
  d={'Vabc':'输入三相相电压。','Iabc':'输入电机三相电流。','line_voltage':'组合 [Vab Vbc Vca]。','Vab':'计算 va−vb。','Vbc':'计算 vb−vc。','Vca':'计算 vc−va。',
  'Vab_Vbc_Vca':'端口 1 输出线电压，V。','Iabc_measured':'端口 2 原样输出输入电流，A。','Vabc_to_motor':'端口 3 原样传递相电压，V。'}
  assert n in d,(parent,n);return d[n]
 if parent in ('Bus Selector','Bus Selector1'):
  if t=='Inport':return '接收 [ia ib ic ωm θm Te vq_actual vd_actual] 八元素向量。'
  if t=='Demux':return '将测量向量按原顺序拆为八个标量。'
  m=re.search(r'(\d+)$',n);assert m,(parent,n)
  index=int(m.group(1)); label=['ia','ib','ic','ωm','θm','Te','vq_actual','vd_actual'][index-1]
  return ('输出' if t=='Outport' else '终止本子系统不用的')+'第 '+str(index)+' 项 '+label+'。'
 if parent=='Subsystem':
  d={'theta_m':'输入机械角度 θm，rad。','Gain1':'机械角乘极对数 p。','Constant':'周期常量 2π。','Mod':'对 p·θm 取 mod 2π，得到 [0,2π) 电角度。','theta_e':'输出 theta_e，rad。','Scope':'观察电角度环绕。'}
  assert n in d,(parent,n);return d[n]
 if parent in ('Chirp1','Chirp2'):
  d={'sweep_time':'10 s 重复斜坡，输出当前扫频周期内的时间。','linear_chirp':'计算 cos(π·'+('500' if parent=='Chirp1' else '5000')+'·t²/10)，构成线性扫频余弦。',
  'sample':'以 dt 采样保持 Chirp。','chirp':'输出单位幅值辅助 Chirp。'}
  assert n in d,(parent,n);return d[n]
 if parent=='Validation logging' and t=='ToWorkspace':return '保存 '+p['VariableName']+' 为 timeseries，SampleTime=dt，Decimation=1，样本上限 inf。'
 if parent=='SVPWM Function v1.0':
  d={'alpha':'端口 1 输入 α 轴电压要求，V。','beta':'端口 2 输入 β 轴电压要求，V。','Vdc':'端口 3 输入母线，V。','Ts':'端口 4 输入 PWM 周期，s。',
  'Tcm':'输出三相比较阈值 [TcmA TcmB TcmC]，s。','sector':'输出符号判定的扇区编码。',
  'SFunction':'Stateflow 生成的执行封装 sf_sfun；用户算法是附录中的 MATLAB Function。',
  'Demux':'Stateflow 自动生成的内部输出拆分，不是额外的 SVPWM 算法。',
  'Terminator':'终止 Stateflow 封装不用的内部输出。'}
  assert n in d,(parent,n);return d[n]
 raise ValueError(('Unexplained block',s,parent,n,t))

def connections(b):
 a=[];t=b['type'];p=b['params']
 for v in arr(b['inputs']):
  source=BY.get(v['source']);assert source
  a.append('入'+str(v['port'])+'←'+sid(source)+'/出'+str(v['source_port']))
 if t=='From':a.append('标签←'+sid(G[p['GotoTag']]))
 for v in arr(b['outputs']):
  dest=[]
  for target in v['destinations']:
   path,port=target.rsplit('#in:',1); dest.append(sid(BY[path])+'/入'+port)
  if dest:a.append('出'+str(v['port'])+'→'+'、'.join(dest))
 if t=='Goto':a.append('标签→'+('、'.join(sid(x) for x in F[p['GotoTag']]) or '无 From'))
 if not a:return '无外部端口；内部通过标签读取日志信号。'
 return '；'.join(a)
def table(blocks):
 rows=['| 块名和 SID | 当前作用及参数 | 直接连接 |','| --- | --- | --- |']
 for b in blocks:
  desc=role(b);assert desc
  rows.append('| '+clean(b['name'])+' ['+sid(b)+'] | '+desc.replace('|','／')+' | '+connections(b)+' |')
 return '\n'.join(rows)

# Crop the real nine panel bounds from distinct background colors.
im=Image.open('D:/simulink/modern/layout/flat_tidy.png').convert('RGB')
colors=sorted(im.getcolors(im.width*im.height),reverse=True)
panels=[]
pixels=np.asarray(im)
def largest_run(indices):
 runs=np.split(indices,np.where(np.diff(indices)>1)[0]+1)
 best=max(runs,key=len);return int(best[0]),int(best[-1])+1
for count,color in colors:
 if count<1000000 or color==(255,255,255):continue
 mask=np.all(pixels==color,axis=2)
 # Some small blocks elsewhere reuse a panel color; global getbbox merges them.
 y1,y2=largest_run(np.flatnonzero(mask.any(axis=1)))
 x1,x2=largest_run(np.flatnonzero(mask[y1:y2].any(axis=0)))
 box=(x1,y1,x2,y2)
 assert abs((x2-x1)/(y2-y1)-860/580)<.01,box
 panels.append((box,color))
assert len(panels)==9,panels
panels.sort(key=lambda x:(round(x[0][1]/100)*100,x[0][0]))
for i,(box,_) in enumerate(panels,1):im.crop(box).save(ASSET/f'region_{i:02}.png')

template=ROOT/'handbook_template.md'
if not template.exists():shutil.copyfile(ROOT/'PMSM模型逐模块解读.md',template)
content=template.read_text(encoding='utf-8')
for i,ids in enumerate(groups,1):content=content.replace(f'<!-- TOP_BLOCKS_{i:02} -->','### 本区小模块索引\n\n'+table([BS[str(x)] for x in ids]))
parents=[]
for ids in groups:
 for s in ids:
  path=BS[str(s)]['path']
  if any(b['parent']==path for b in B):parents.append(path)
internal='表中 SID 省略共同前缀 PMSM_FOC_R2026a。连接写法“入2←190/出1”表示该块第 2 输入来自 SID 190 的第 1 输出；“出1→155/入1”表示输出送往对应输入。“标签”表示 From 和 Goto 的隐式连接。内部名称按当前文件保留。\n\n'
for parent in parents:
 blocks=[b for b in B if b['parent']==parent]
 internal+='### '+clean(parent.removeprefix(M+'/'))+' 内部块\n\n'+table(blocks)+'\n\n'
content=content.replace('<!-- INTERNAL_BLOCKS -->',internal.strip())
functions=arr(E['matlab_functions']);assert len(functions)==1
source=functions[0]['script']
code='\n'.join(line.split('%')[0].rstrip() if '\ufffd' in line else line for line in source.splitlines())
code=re.sub(r'\n{3,}','\n\n',code)
content=content.replace('以下代码来自当前模型内嵌 MATLAB Function，保留现状以便阅读。','以下代码来自当前模型内嵌 MATLAB Function，保留原语句和可正常显示的注释，供逐行对照。')
content=content.replace('<!-- SVPWM_SOURCE -->','```matlab\n'+code+'\n```')
assert '<!--' not in content and '\ufffd' not in content
content=content.replace('图示文件位于本目录', '索引连接以 SID 表示，例如“入2←190/出1”表示该块第 2 输入来自 SID 190 的第 1 输出，“标签”表示 From/Goto 连接。SID 省略共同前缀 PMSM_FOC_R2026a。可在 MATLAB 用 Simulink.ID.getHandle 定位 SID，再用 hilite_system 高亮对应块。\n\n本文逐项登记当前文件中导出的显式块。Carrier Wave 和 sweep_time 等库块按接口及参数解释，不把整个 Simulink 库的隐藏实现计入 329 个块。\n\n图示文件位于本目录')
(ROOT/'PMSM模型逐模块解读.md').write_text(content,encoding='utf-8')

# Native OMML equations, using the small supported subset below.
# Every operator used by the source is explicitly recognized and later rendered for QA.
def me(tag,children=(),text=None):
 e=OxmlElement('m:'+tag)
 if text is not None:e.text=text
 for c in children:e.append(c)
 return e
def mr(text):
 r=me('r');pr=me('rPr');sty=me('sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
 props=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),'Cambria Math');f.set(qn('w:hAnsi'),'Cambria Math');props.append(f);r.append(props)
 t=me('t',text=text);t.set(qn('xml:space'),'preserve');r.append(t);return r
GREEK={'psi':'ψ','omega':'ω','theta':'θ','alpha':'α','beta':'β','pi':'π','Delta':'Δ'}
class MathParser:
 def __init__(self,s):self.s=s;self.i=0
 def row(self,end=None):
  out=[]
  while self.i<len(self.s) and (end is None or self.s[self.i]!=end):
   if self.s[self.i].isspace():self.i+=1;continue
   a=self.atom();sub=sup=None
   while self.i<len(self.s) and self.s[self.i] in '_^':
    sym=self.s[self.i];self.i+=1;v=self.arg()
    if sym=='_':sub=v
    else:sup=v
   if sub is not None and sup is not None:a=me('sSubSup',[me('e',[a]),me('sub',sub),me('sup',sup)])
   elif sub is not None:a=me('sSub',[me('e',[a]),me('sub',sub)])
   elif sup is not None:a=me('sSup',[me('e',[a]),me('sup',sup)])
   out.append(a)
  if end is not None:
   assert self.i<len(self.s) and self.s[self.i]==end,(self.s,self.i,end);self.i+=1
  return out
 def arg(self):
  while self.i<len(self.s) and self.s[self.i].isspace():self.i+=1
  if self.s[self.i]=='{':self.i+=1;return self.row('}')
  return [self.atom()]
 def atom(self):
  c=self.s[self.i];self.i+=1
  if c=='{':return me('box',[me('e',self.row('}'))])
  if c=='\\':
   match=re.match(r'[A-Za-z]+',self.s[self.i:])
   if not match:
    c=self.s[self.i];self.i+=1;return mr(' ' if c in ',;!' else c)
   cmd=match.group();self.i+=len(cmd)
   if cmd in GREEK:return mr(GREEK[cmd])
   if cmd=='frac':return me('f',[me('num',self.arg()),me('den',self.arg())])
   if cmd=='sqrt':return me('rad',[me('radPr',[self.attr('degHide','1')]),me('deg'),me('e',self.arg())])
   if cmd=='dot':return me('acc',[me('accPr',[self.attr('chr','̇')]),me('e',self.arg())])
   if cmd in ['cos','sin']:return mr(cmd+' ')
   if cmd=='operatorname':return me('box',[me('e',self.arg())])
   if cmd in ['quad','qquad']:return mr('  ')
   if cmd=='leq':return mr('≤')
   raise ValueError(('Unsupported math command',cmd,self.s))
  return mr('−' if c=='-' else c)
 def attr(self,t,val):
  e=me(t);e.set(qn('m:val'),val);return e

def setfont(r,size=None,bold=None,font='Microsoft YaHei'):
 r.font.name=font
 pr=r._element.get_or_add_rPr();f=pr.find(qn('w:rFonts'))
 if f is None:f=OxmlElement('w:rFonts');pr.insert(0,f)
 for k in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+k),font)
 if size:r.font.size=Pt(size)
 if bold is not None:r.bold=bold

D=Document();sec=D.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.72)
sec.header_distance=sec.footer_distance=Inches(.3)
for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 style=D.styles[st];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
 style._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
 f=style._element.rPr.findall(qn('w:rFonts'))[-1]
 for k in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+k),'Microsoft YaHei')
normal=D.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.22;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',25),('Subtitle',12),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
 st=D.styles[name];st.font.size=Pt(size);st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.keep_with_next=True
 if name!='Subtitle':st.font.bold=True
header=sec.header.paragraphs[0];header.text='PMSM FOC 模型解读与 STM32G474 开发计划';setfont(header.runs[0],9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('第 ');setfont(run,9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
setfont(footer.add_run(' 页'),9)

def hyperlink(p,label,url):
 rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'245A81');pr.append(co);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);el.append(r);p._p.append(el)
def inline(p,s,size=None):
 parts=re.split(r'(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)',s)
 for x in parts:
  if not x:continue
  m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',x)
  if m:hyperlink(p,m[1],m[2]);continue
  bold=x.startswith('**');code=x.startswith('`');x=x[2:-2] if bold else x[1:-1] if code else x
  # Long paths/hashes can otherwise overflow a narrow cell.
  if code and len(x)>45:x='\u200b'.join(x[i:i+18] for i in range(0,len(x),18))
  r=p.add_run(x);setfont(r,size,bold)
def titletext(s):return re.sub(r'[^\w\s\u4e00-\u9fff]+',' ',s).replace('_',' ').strip()
def addtable(rows):
 n=len(rows[0]);t=D.add_table(rows=0,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 total=7.06
 widths=([1.7,3.15,2.21] if n==3 else [1.18,1.6,1.6,2.68] if n==4 else [total/n]*n)
 if n==9:widths=[.63]+[.80375]*8
 if n==2:widths=[1.8,5.26]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 for i,row in enumerate(rows):
  rr=t.add_row()
  # Keep each individual record intact; the table can paginate with repeated headings.
  trpr=rr._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,(cell,txt,w) in enumerate(zip(rr.cells,row,widths)):
   cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
   for side,val in [('top',75),('bottom',75),('left',95),('right',95)]:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');margin.append(x)
   cp.append(margin)
   if i==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E2EAF0');cp.append(fill)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
   if n==9:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   inline(p,txt,9.5 if n<9 else 8.5)
   if i==0:
    for r in p.runs:r.bold=True
 # Subsequent narrative paragraphs get spacing before them; no empty spacer row.
def render_md(md,part):
 lines=md.splitlines();i=0
 while i<len(lines):
  s=lines[i].strip();i+=1
  if not s:continue
  if s.startswith('# '):continue
  if s.startswith('```'):
   code=[]
   while i<len(lines) and not lines[i].strip().startswith('```'):code.append(lines[i]);i+=1
   i+=1
   for line in code:
    p=D.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(10.5)
    r=p.add_run(line if line else ' ');setfont(r,8.5,font='Consolas')
   continue
  if s.startswith('$$'):
   eq=s[2:-2];p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   math=me('oMathPara',[me('oMath',MathParser(eq).row())]);p._p.append(math);continue
  if s.startswith('|'):
   rows=[]
   while True:
    if not re.match(r'^\|\s*:?-',s):rows.append([x.strip() for x in s.strip('|').split('|')])
    if i>=len(lines) or not lines[i].strip().startswith('|'):break
    s=lines[i].strip();i+=1
   assert all(len(r)==len(rows[0]) for r in rows),rows
   addtable(rows);continue
  if s.startswith('!['):
   m=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',s);assert m
   p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   p.add_run().add_picture(str(ROOT/m[2]),width=Inches(6.9))
   p.paragraph_format.keep_with_next=True
   cap=D.add_paragraph('当前模型分区图 '+m[1]);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for r in cap.runs:setfont(r,9)
   continue
  if s.startswith('##'):
   level=len(s)-len(s.lstrip('#'))-1;level=min(level,3);text=s.lstrip('#').strip()
   p=D.add_heading(titletext(text),level=level)
   # Begin each main control region on a new page; other subchapters flow naturally.
   if part==1 and re.match(r'^0[1-9] ',text):p.paragraph_format.page_break_before=True
   if part==1 and text=='附录 当前 SVPWM 函数':p.paragraph_format.page_break_before=True
   continue
  p=D.add_paragraph()
  prev=p._p.getprevious()
  if prev is not None and prev.tag==qn('w:tbl'):p.paragraph_format.space_before=Pt(6)
  inline(p,s)

p=D.add_paragraph('PMSM FOC 模型解读\n与 STM32G474 开发计划',style='Title')
D.add_paragraph('基于当前平铺模型的学习与开发手册',style='Subtitle')
D.add_paragraph('MATLAB R2026a Update 5   STM32G474RET6\n2026 年 10 月 11 日')
D.add_paragraph('本手册把当前模型的实现与后续修改建议分开说明。第一部分按九个色域解读，并登记全部 329 个显式块；第二部分给出用真实电路、采样时序和固件逐步校准仿真的开发计划。')
D.add_paragraph('建议先看控制链和速度环，再依次进入电流环、坐标变换、SVPWM、电机及反馈。逐块索引适合查阅，不必在第一次阅读时逐行通读。')
D.add_heading('内容导览',level=1)
for item in ['第一部分 模型逐模块解读','01 速度环  02 电流环  03 坐标变换','04 SVPWM  05 逆变器与 PMSM  06 采样反馈','07 数据观测  08 调制观测  09 辅助观测','已有验证和学习顺序','内部 165 个小模块索引与当前 SVPWM 源码','第二部分 STM32G474 仿真开发计划','算法修复  真实参数  采样时序  测量链','功率级与电机  C 固件  处理器验证  实物校准']:
 D.add_paragraph(item)
D.add_page_break();D.add_heading('第一部分 模型逐模块解读',level=1);render_md(content,1)
D.add_page_break();D.add_heading('第二部分 STM32G474 仿真开发计划',level=1)
plan=(ROOT/'STM32G474仿真开发计划.md').read_text(encoding='utf-8');render_md(plan,2)
# The default DOCX template may contain a theme-colored title paragraph border.
for element in [D.styles.element,D.element]:
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
D.styles['Subtitle'].font.italic=False
out=ROOT/'PMSM模型解读与STM32G474开发计划.docx';D.save(out)
assert hashlib.sha256(MODEL.read_bytes()).hexdigest()==EXPECTED
covered=[b for ids in groups for b in [BS[str(s)] for s in ids]]+[b for b in B if b['parent']!=M]
assert len(covered)==329 and len({b['sid'] for b in covered})==329
audit={'model_unchanged_sha256':EXPECTED,'top_level_count':164,'internal_count':165,'total_count':329,
'region_counts':[len(x) for x in groups],'blocks':[{'sid':b['sid'],'name':clean(b['name']),'description':role(b)} for b in covered],
'native_equation_count':len(re.findall(r'^\$\$',content,re.M)),'docx':str(out),'source_evidence':str(ROOT/'current_model_evidence.json')}
(ROOT/'handbook_coverage.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print('BUILT',out,'ROOT',164,'INTERNAL',165,'EQUATIONS',audit['native_equation_count'])
